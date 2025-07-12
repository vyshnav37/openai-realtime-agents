from flask import Flask, request, jsonify
from flask_cors import CORS
from flask_sqlalchemy import SQLAlchemy
from datetime import datetime, timedelta
import re
import os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
import sqlite3
import json

app = Flask(__name__)
CORS(app)

# Database configuration
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///pool_maintenance.db'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
db = SQLAlchemy(app)

# Villa patterns as specified
VILLA_PATTERNS = [
    '101-112', '201-207', '301-308', 
    '401-409', '501-511', '601-610',
    '701-709', '801-806', '901-907'
]

# Generate all villa numbers from patterns
def generate_villa_numbers():
    villas = []
    for pattern in VILLA_PATTERNS:
        start, end = pattern.split('-')
        for i in range(int(start), int(end) + 1):
            villas.append(str(i))
    return villas

ALL_VILLAS = generate_villa_numbers()

# Database Models
class Villa(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    villa_number = db.Column(db.String(10), unique=True, nullable=False)
    series = db.Column(db.String(10), nullable=False)
    status = db.Column(db.String(10), default='VD')  # VD = Vacant Dirty, OCC = Occupied
    
class MaintenanceRecord(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    villa_id = db.Column(db.Integer, db.ForeignKey('villa.id'), nullable=False)
    timestamp = db.Column(db.DateTime, default=datetime.utcnow)
    chlorine_level = db.Column(db.Float)
    ph_level = db.Column(db.Float)
    temperature = db.Column(db.Float)
    status = db.Column(db.String(50))
    notes = db.Column(db.Text)
    staff_name = db.Column(db.String(100))
    time_in = db.Column(db.DateTime)
    time_out = db.Column(db.DateTime)
    
class FountainMaintenance(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    villa_id = db.Column(db.Integer, db.ForeignKey('villa.id'), nullable=False)
    timestamp = db.Column(db.DateTime, default=datetime.utcnow)
    fountain_name = db.Column(db.String(100))
    issue_description = db.Column(db.Text)
    status = db.Column(db.String(50))
    resolved = db.Column(db.Boolean, default=False)

# Message parsing utilities
class MessageParser:
    @staticmethod
    def parse_maintenance_message(message):
        """Parse WhatsApp message for maintenance data"""
        data = {
            'villa_number': None,
            'chlorine_level': None,
            'ph_level': None,
            'temperature': None,
            'status': 'VD',
            'notes': '',
            'staff_name': '',
            'timestamp': datetime.now()
        }
        
        # Extract villa number
        villa_match = re.search(r'\b(1[0-1][0-9]|2[0-0][0-7]|3[0-0][0-8]|4[0-0][0-9]|5[0-1][0-9]|6[0-1][0-9]|7[0-0][0-9]|8[0-0][0-6]|9[0-0][0-7])\b', message)
        if villa_match:
            data['villa_number'] = villa_match.group(1)
        
        # Extract chlorine level
        chlorine_match = re.search(r'chlorine\s*(\d+\.?\d*)', message, re.IGNORECASE)
        if chlorine_match:
            data['chlorine_level'] = float(chlorine_match.group(1))
        
        # Extract pH level
        ph_match = re.search(r'ph\s*(\d+\.?\d*)', message, re.IGNORECASE)
        if ph_match:
            data['ph_level'] = float(ph_match.group(1))
        
        # Extract temperature
        temp_match = re.search(r'temp\w*\s*(\d+\.?\d*)', message, re.IGNORECASE)
        if temp_match:
            data['temperature'] = float(temp_match.group(1))
        
        # Determine status
        if 'occupied' in message.lower():
            data['status'] = 'OCC'
        elif 'vacant' in message.lower():
            data['status'] = 'VD'
        elif 'clean' in message.lower():
            data['status'] = 'VC'
        
        # Extract notes
        data['notes'] = message
        
        return data
    
    @staticmethod
    def parse_fountain_message(message):
        """Parse fountain maintenance messages"""
        data = {
            'villa_number': None,
            'fountain_name': 'Main Fountain',
            'issue_description': message,
            'status': 'Reported',
            'resolved': False
        }
        
        # Extract villa number
        villa_match = re.search(r'\b(1[0-1][0-9]|2[0-0][0-7]|3[0-0][0-8]|4[0-0][0-9]|5[0-1][0-9]|6[0-1][0-9]|7[0-0][0-9]|8[0-0][0-6]|9[0-0][0-7])\b', message)
        if villa_match:
            data['villa_number'] = villa_match.group(1)
        
        # Determine fountain name by series
        if data['villa_number']:
            series = data['villa_number'][0]
            data['fountain_name'] = f'Series{series}-MainFountain'
        
        return data

# Excel utilities
class ExcelManager:
    @staticmethod
    def create_maintenance_template():
        """Create Excel template for maintenance records"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Pool Maintenance"
        
        # Headers
        headers = ['Villa No', 'Date', 'Time In', 'Time Out', 'Staff', 'Status', 'Temperature', 'Chlorine (ppm)', 'pH Level', 'Notes']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col)
            cell.value = header
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")
        
        return wb
    
    @staticmethod
    def fill_maintenance_data(records):
        """Fill Excel with maintenance data"""
        wb = ExcelManager.create_maintenance_template()
        ws = wb.active
        
        row = 2
        for record in records:
            ws.cell(row=row, column=1, value=record.villa.villa_number)
            ws.cell(row=row, column=2, value=record.timestamp.strftime('%Y-%m-%d'))
            ws.cell(row=row, column=3, value=record.time_in.strftime('%H:%M') if record.time_in else '')
            ws.cell(row=row, column=4, value=record.time_out.strftime('%H:%M') if record.time_out else '')
            ws.cell(row=row, column=5, value=record.staff_name)
            ws.cell(row=row, column=6, value=record.status)
            ws.cell(row=row, column=7, value=record.temperature)
            ws.cell(row=row, column=8, value=record.chlorine_level)
            ws.cell(row=row, column=9, value=record.ph_level)
            ws.cell(row=row, column=10, value=record.notes)
            row += 1
        
        return wb

# Routes
@app.route('/api/villas', methods=['GET'])
def get_villas():
    """Get all villas with their current status"""
    villas = Villa.query.all()
    return jsonify([{
        'id': villa.id,
        'villa_number': villa.villa_number,
        'series': villa.series,
        'status': villa.status
    } for villa in villas])

@app.route('/api/process-message', methods=['POST'])
def process_message():
    """Process WhatsApp message for maintenance data"""
    data = request.json
    message = data.get('message', '')
    
    # Determine if it's a fountain or pool maintenance message
    if any(keyword in message.lower() for keyword in ['fountain', 'nozzle', 'pump', 'filter']):
        fountain_data = MessageParser.parse_fountain_message(message)
        
        if fountain_data['villa_number']:
            villa = Villa.query.filter_by(villa_number=fountain_data['villa_number']).first()
            if villa:
                fountain_record = FountainMaintenance(
                    villa_id=villa.id,
                    fountain_name=fountain_data['fountain_name'],
                    issue_description=fountain_data['issue_description'],
                    status=fountain_data['status']
                )
                db.session.add(fountain_record)
                db.session.commit()
                
                return jsonify({
                    'success': True,
                    'type': 'fountain',
                    'villa_number': fountain_data['villa_number'],
                    'message': 'Fountain issue recorded'
                })
    else:
        # Process as pool maintenance
        pool_data = MessageParser.parse_maintenance_message(message)
        
        if pool_data['villa_number']:
            villa = Villa.query.filter_by(villa_number=pool_data['villa_number']).first()
            if villa:
                maintenance_record = MaintenanceRecord(
                    villa_id=villa.id,
                    chlorine_level=pool_data['chlorine_level'],
                    ph_level=pool_data['ph_level'],
                    temperature=pool_data['temperature'],
                    status=pool_data['status'],
                    notes=pool_data['notes'],
                    time_in=datetime.now()
                )
                db.session.add(maintenance_record)
                db.session.commit()
                
                return jsonify({
                    'success': True,
                    'type': 'pool',
                    'villa_number': pool_data['villa_number'],
                    'data': pool_data
                })
    
    return jsonify({'success': False, 'message': 'Could not parse message'})

@app.route('/api/maintenance-records', methods=['GET'])
def get_maintenance_records():
    """Get maintenance records with optional filtering"""
    villa_id = request.args.get('villa_id')
    date_from = request.args.get('date_from')
    date_to = request.args.get('date_to')
    
    query = MaintenanceRecord.query.join(Villa)
    
    if villa_id:
        query = query.filter(MaintenanceRecord.villa_id == villa_id)
    
    if date_from:
        query = query.filter(MaintenanceRecord.timestamp >= datetime.strptime(date_from, '%Y-%m-%d'))
    
    if date_to:
        query = query.filter(MaintenanceRecord.timestamp <= datetime.strptime(date_to, '%Y-%m-%d'))
    
    records = query.all()
    
    return jsonify([{
        'id': record.id,
        'villa_number': record.villa.villa_number,
        'timestamp': record.timestamp.isoformat(),
        'chlorine_level': record.chlorine_level,
        'ph_level': record.ph_level,
        'temperature': record.temperature,
        'status': record.status,
        'notes': record.notes,
        'staff_name': record.staff_name
    } for record in records])

@app.route('/api/export-excel', methods=['GET'])
def export_excel():
    """Export maintenance records to Excel"""
    date_from = request.args.get('date_from')
    date_to = request.args.get('date_to')
    
    query = MaintenanceRecord.query.join(Villa)
    
    if date_from:
        query = query.filter(MaintenanceRecord.timestamp >= datetime.strptime(date_from, '%Y-%m-%d'))
    
    if date_to:
        query = query.filter(MaintenanceRecord.timestamp <= datetime.strptime(date_to, '%Y-%m-%d'))
    
    records = query.all()
    wb = ExcelManager.fill_maintenance_data(records)
    
    # Save to file
    filename = f"pool_maintenance_{datetime.now().strftime('%Y%m%d')}.xlsx"
    wb.save(filename)
    
    return jsonify({
        'success': True,
        'filename': filename,
        'records_count': len(records)
    })

@app.route('/api/dashboard-stats', methods=['GET'])
def get_dashboard_stats():
    """Get dashboard statistics"""
    total_villas = Villa.query.count()
    total_records = MaintenanceRecord.query.count()
    
    # Get today's maintenance count
    today = datetime.now().date()
    today_records = MaintenanceRecord.query.filter(
        MaintenanceRecord.timestamp >= today,
        MaintenanceRecord.timestamp < today + timedelta(days=1)
    ).count()
    
    # Get villas with low pH (< 7.2)
    low_ph_records = MaintenanceRecord.query.filter(
        MaintenanceRecord.ph_level < 7.2,
        MaintenanceRecord.timestamp >= today - timedelta(days=1)
    ).count()
    
    # Get fountain issues
    fountain_issues = FountainMaintenance.query.filter_by(resolved=False).count()
    
    return jsonify({
        'total_villas': total_villas,
        'total_records': total_records,
        'today_records': today_records,
        'low_ph_alerts': low_ph_records,
        'fountain_issues': fountain_issues
    })

# Initialize database
def init_db():
    """Initialize database with villa data"""
    with app.app_context():
        db.create_all()
        
        # Check if villas are already populated
        if Villa.query.count() == 0:
            # Populate villas
            for villa_num in ALL_VILLAS:
                series = villa_num[0]  # First digit indicates series
                villa = Villa(
                    villa_number=villa_num,
                    series=f'Series{series}',
                    status='VD'
                )
                db.session.add(villa)
            
            db.session.commit()
            print(f"Initialized {len(ALL_VILLAS)} villas")

if __name__ == '__main__':
    init_db()
    app.run(debug=True, port=5000)